"""Migration regressions use an isolated SQLite DB; no production data or paid calls."""
import io
import unittest
from datetime import datetime, timedelta
from unittest.mock import patch

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from main import app
from database import Base, get_db
from models import AdminUser, Document, Generation, Tenant
from routers.admin import _create_admin_token
from services.knowledge import knowledge_service


class MigrationTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine)

        def test_db():
            with self.Session() as db:
                yield db
        app.dependency_overrides[get_db] = test_db
        self.addCleanup(app.dependency_overrides.clear)
        self.addCleanup(self.engine.dispose)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)
        self.mail = patch('routers.auth.is_email_service_configured', return_value=False)
        self.mail.start()
        self.addCleanup(self.mail.stop)
        self.index = patch.object(knowledge_service, '_ensure_initialized', side_effect=RuntimeError('index offline'))
        self.index.start()
        self.addCleanup(self.index.stop)

    def register(self, name):
        response = self.client.post('/api/auth/register', json={
            'email': f'{name}@example.test', 'password': 'Migration123!', 'nickname': name,
        })
        self.assertEqual(response.status_code, 200, response.text)
        return {'Authorization': 'Bearer ' + response.json()['token']}

    def test_real_registration_refresh_password_and_tenant_isolation(self):
        self.assertEqual(self.client.get('/api/documents').status_code, 401)
        alice, bob = self.register('alice'), self.register('bob')
        response = self.client.post('/api/documents/add-text', headers=alice, json={'name': 'Private', 'content': 'Only Alice can read this saved product document.'})
        self.assertEqual(response.status_code, 200, response.text)
        doc_id = response.json()['id']
        self.assertEqual(self.client.get(f'/api/documents/{doc_id}', headers=bob).status_code, 404)
        self.assertEqual(self.client.delete(f'/api/documents/{doc_id}', headers=bob).status_code, 404)
        self.assertEqual(self.client.get('/api/documents', headers=bob).json(), [])
        refreshed = self.client.post('/api/auth/refresh', headers=alice)
        self.assertEqual(refreshed.status_code, 200)
        self.assertEqual(self.client.post('/api/auth/change-password', headers=alice, json={'old_password': 'Migration123!', 'new_password': 'Changed456!'}).status_code, 200)
        self.assertEqual(self.client.post('/api/auth/login', json={'email': 'alice@example.test', 'password': 'Migration123!'}).status_code, 401)
        login = self.client.post('/api/auth/login', json={'email': 'alice@example.test', 'password': 'Changed456!'})
        self.assertEqual(login.status_code, 200)
        self.assertEqual(self.client.get('/api/auth/me', headers=alice).status_code, 401)
        with TestClient(app) as reopened:
            response = reopened.get(f'/api/documents/{doc_id}', headers={'Authorization': 'Bearer ' + login.json()['token']})
            self.assertEqual(response.status_code, 200)
            self.assertIn('Only Alice', response.json()['content'])

    def test_pdf_docx_and_txt_parse_without_vector_service(self):
        import fitz
        from docx import Document as WordDocument
        import tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pdf = fitz.open()
            pdf.new_page().insert_text((40, 40), 'Actual PDF product material')
            pdf.save(root / 'sample.pdf')
            pdf.close()
            word = WordDocument()
            word.add_paragraph('Actual Word product material')
            word.save(root / 'sample.docx')
            (root / 'sample.txt').write_text('Actual TXT product material', encoding='utf-8')
            with self.Session() as db:
                for ext in ['pdf', 'docx', 'txt']:
                    with self.subTest(ext=ext):
                        doc = Document(name=ext, file_type=ext, file_path=str(root / f'sample.{ext}'))
                        db.add(doc); db.commit()
                        self.assertEqual(knowledge_service.process_document(db, doc.id), 0)
                        db.refresh(doc)
                        self.assertIn('product material', doc.content)
                        self.assertGreater(doc.chunk_count, 0)
                        self.assertFalse(doc.indexed)

    def test_admin_analytics_uses_saved_data_and_filters_dates_and_platforms(self):
        self.assertEqual(self.client.get('/api/admin/generations/analytics').status_code, 401)
        user = self.register('ordinary')
        self.assertEqual(self.client.get('/api/admin/generations/analytics', headers=user).status_code, 401)
        with self.Session() as db:
            admin = AdminUser(username='test-admin', password_hash='not-used', is_active=True)
            tenant = Tenant(name='Analytics tenant')
            db.add_all([admin, tenant]); db.commit()
            headers = {'Authorization': 'Bearer ' + _create_admin_token(admin.id)}
            db.add_all([
                Generation(tenant_id=tenant.id, topic='A', platform='douyin', output_full='12345', rating=4),
                Generation(tenant_id=tenant.id, topic='B', platform='xiaohongshu', output_full='1234567'),
                Generation(tenant_id=tenant.id, topic='Old', output_full='old', created_at=datetime.utcnow() - timedelta(days=40)),
            ]); db.commit()
        result = self.client.get('/api/admin/generations/analytics?time_range=7d', headers=headers)
        self.assertEqual(result.status_code, 200, result.text)
        data = result.json()
        self.assertEqual(data['summary']['total_generations'], 2)
        self.assertEqual(data['summary']['avg_word_count'], 6)
        self.assertEqual(data['summary']['avg_rating'], 4)
        self.assertEqual(data['summary']['rated_count'], 1)
        self.assertEqual(sum(d['count'] for d in data['trend_days']), 2)
        self.assertEqual(self.client.get('/api/admin/generations/analytics?time_range=bad', headers=headers).status_code, 422)
        listing = self.client.get('/api/admin/generations?platform=douyin', headers=headers).json()
        self.assertEqual(listing['total'], 1)
        self.assertEqual(listing['items'][0]['word_count'], 5)
        with self.Session() as db:
            db.query(Generation).delete(); db.commit()
        empty = self.client.get('/api/admin/generations/analytics', headers=headers).json()
        self.assertEqual(empty['summary']['total_generations'], 0)
        self.assertIsNone(empty['summary']['avg_rating'])
        self.assertEqual(empty['platforms'], [])
